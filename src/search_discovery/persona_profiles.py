import argparse
import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

from src.search_discovery.io import write_json, write_jsonl


SCHEMA_VERSION = "0.1"
REQUIRED_HEADERS = ("一级赛道", "二级赛道", "人设")

_PROFILE_TYPE_RULES = (
    ("tech_ai_creator", ("科技", "AI", "人工智能", "工具", "开发", "编程", "效率")),
    ("business_startup_creator", ("财经", "商业", "理财", "基金", "股票", "消费", "投资", "职场")),
)

_CONTENT_MODE_RULES = (
    ("工具测评", ("工具", "体验", "测试", "测评", "应用")),
    ("教程实践", ("教程", "实践", "入门", "学习", "备考", "科研", "汇报", "写作")),
    ("案例拆解", ("案例", "拆解", "复盘", "分析")),
    ("避坑清单", ("避坑", "风险", "不讲", "普通人")),
    ("趋势观察", ("观察", "趋势", "城市", "市场", "行业")),
    ("新闻解读", ("新闻", "热点", "政策", "回应", "事件")),
)

_KEYWORD_VOCABULARY = (
    "AI",
    "AI工具",
    "写作",
    "学习",
    "汇报",
    "科研",
    "论文",
    "备考",
    "简历",
    "面试",
    "校招",
    "实习",
    "基金",
    "存款",
    "消费",
    "风险",
    "理财",
    "生活化案例",
    "城市",
    "博物馆",
    "建筑",
    "商业",
    "职场",
    "效率",
    "工具",
    "测评",
    "教程",
    "案例",
    "拆解",
)

_STOP_WORDS = {
    "一个",
    "一种",
    "专注",
    "专门",
    "分享",
    "视角",
    "背景",
    "故事",
    "场景",
    "经验",
    "普通人",
}


@dataclass(frozen=True)
class PersonaProfile:
    profile_id: str
    primary_track: str
    secondary_track: str
    persona_text: str
    profile_type: str
    track_tags: list[str] = field(default_factory=list)
    custom_keywords: list[str] = field(default_factory=list)
    content_modes: list[str] = field(default_factory=list)
    retrieval_text: str = ""
    sensitivity_level: str = "internal"
    source: dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> dict[str, Any]:
        row = asdict(self)
        row["schema_version"] = SCHEMA_VERSION
        return row


def load_persona_profiles_from_excel(path: Path) -> list[PersonaProfile]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 .xlsx 需要安装 openpyxl；请先运行项目依赖同步。") from exc

    workbook = load_workbook(path, data_only=True, read_only=True)
    personas: list[PersonaProfile] = []
    counter = 1
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        header_index = _find_header_row(rows)
        if header_index is None:
            continue
        header = [str(value).strip() if value is not None else "" for value in rows[header_index]]
        indexes = _header_indexes(header)
        for row_offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
            primary_track = _cell(row, indexes["一级赛道"])
            secondary_track = _cell(row, indexes["二级赛道"])
            persona_text = _cell(row, indexes["人设"])
            if not any((primary_track, secondary_track, persona_text)):
                continue
            personas.append(
                build_persona_profile(
                    profile_id=f"persona_{counter:04d}",
                    primary_track=primary_track,
                    secondary_track=secondary_track,
                    persona_text=persona_text,
                    source={"file": path.name, "sheet": sheet.title, "row": row_offset},
                )
            )
            counter += 1
    return personas


def build_persona_profile(
    *,
    profile_id: str,
    primary_track: str,
    secondary_track: str,
    persona_text: str,
    source: dict[str, Any],
) -> PersonaProfile:
    track_tags = _unique_nonempty([primary_track, secondary_track])
    profile_type = infer_profile_type(primary_track, secondary_track, persona_text)
    custom_keywords = extract_keywords([primary_track, secondary_track, persona_text])
    content_modes = infer_content_modes(primary_track, secondary_track, persona_text)
    retrieval_text = " ".join(_unique_nonempty([*track_tags, persona_text, *custom_keywords, *content_modes]))
    return PersonaProfile(
        profile_id=profile_id,
        primary_track=primary_track,
        secondary_track=secondary_track,
        persona_text=persona_text,
        profile_type=profile_type,
        track_tags=track_tags,
        custom_keywords=custom_keywords,
        content_modes=content_modes,
        retrieval_text=retrieval_text,
        source=source,
    )


def persona_to_retrieval_document(persona: PersonaProfile) -> dict[str, Any]:
    return persona.to_dict()


def build_creator_profile(persona: PersonaProfile) -> dict[str, Any]:
    return {
        "creator_id": persona.profile_id,
        "role": persona.persona_text,
        "profile_type": persona.profile_type,
        "track_tags": persona.track_tags,
        "custom_keywords": persona.custom_keywords,
        "content_modes": persona.content_modes,
        "platforms": [],
        "content_goal": "基于该用户画像发现适合内容创作的国内热点",
        "exclude_keywords": [],
    }


def write_persona_outputs(personas: list[PersonaProfile], output_dir: Path) -> dict[str, Path]:
    persona_rows = [persona_to_retrieval_document(persona) for persona in personas]
    creator_profiles = [build_creator_profile(persona) for persona in personas]
    paths = {
        "persona_profiles": output_dir / "persona_profiles.jsonl",
        "creator_profiles": output_dir / "creator_profiles.jsonl",
        "index": output_dir / "persona_index.json",
    }
    write_jsonl(paths["persona_profiles"], persona_rows)
    write_jsonl(paths["creator_profiles"], creator_profiles)
    write_json(
        paths["index"],
        {
            "schema_version": SCHEMA_VERSION,
            "total_profiles": len(personas),
            "profiles": [
                {
                    "profile_id": persona.profile_id,
                    "primary_track": persona.primary_track,
                    "secondary_track": persona.secondary_track,
                    "profile_type": persona.profile_type,
                    "keywords_count": len(persona.custom_keywords),
                    "source": persona.source,
                }
                for persona in personas
            ],
        },
    )
    return paths


def infer_profile_type(*values: str) -> str:
    text = " ".join(values)
    for profile_type, markers in _PROFILE_TYPE_RULES:
        if any(marker in text for marker in markers):
            return profile_type
    return "general_hot_topic_creator"


def infer_content_modes(*values: str) -> list[str]:
    text = " ".join(values)
    modes = [mode for mode, markers in _CONTENT_MODE_RULES if any(marker in text for marker in markers)]
    return modes or ["热点观察"]


def extract_keywords(values: Iterable[str], limit: int = 12) -> list[str]:
    text = " ".join(value for value in values if value)
    keywords = [word for word in _KEYWORD_VOCABULARY if word in text]
    for token in re.split(r"[，,。；;、\s]+", text):
        cleaned = token.strip("：:“”\"'（）()")
        if 2 <= len(cleaned) <= 8 and cleaned not in _STOP_WORDS:
            keywords.append(cleaned)
    return _unique_nonempty(keywords)[:limit]


def _find_header_row(rows: list[tuple[Any, ...]]) -> int | None:
    for index, row in enumerate(rows[:20]):
        values = {str(value).strip() for value in row if value is not None}
        if all(header in values for header in REQUIRED_HEADERS):
            return index
    return None


def _header_indexes(header: list[str]) -> dict[str, int]:
    indexes = {name: header.index(name) for name in REQUIRED_HEADERS if name in header}
    missing = [name for name in REQUIRED_HEADERS if name not in indexes]
    if missing:
        raise ValueError(f"Excel 缺少必需表头：{', '.join(missing)}")
    return indexes


def _cell(row: tuple[Any, ...], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _unique_nonempty(values: Iterable[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        cleaned = str(value).strip()
        if not cleaned or cleaned in seen:
            continue
        seen.add(cleaned)
        result.append(cleaned)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert persona Excel rows into retrieval-ready JSONL files.")
    parser.add_argument("--input", required=True, help="Path to Excel file with 一级赛道/二级赛道/人设 columns.")
    parser.add_argument(
        "--output-dir",
        default="data/search_discovery/personas",
        help="Directory for generated persona_profiles.jsonl and creator_profiles.jsonl.",
    )
    args = parser.parse_args()
    personas = load_persona_profiles_from_excel(Path(args.input))
    paths = write_persona_outputs(personas, Path(args.output_dir))
    print(
        json.dumps(
            {
                "profiles_count": len(personas),
                "persona_profiles": paths["persona_profiles"].as_posix(),
                "creator_profiles": paths["creator_profiles"].as_posix(),
                "index": paths["index"].as_posix(),
            },
            ensure_ascii=False,
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
