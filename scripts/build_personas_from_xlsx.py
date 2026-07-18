"""Convert persona xlsx rows into config/profiles/{user_id}.json (v2 schema).

Excel 列： 一级赛道 / 二级赛道 / 人设

Usage:
    python scripts/build_personas_from_xlsx.py --input 人设数据.xlsx
    python scripts/build_personas_from_xlsx.py --input 人设数据.xlsx --use-llm
    python scripts/build_personas_from_xlsx.py --input 人设数据.xlsx --regenerate

Default behavior uses LLM to structure persona and extract keywords.
Falls back to heuristic if LLM is unavailable.

The script deduplicates by `(slug, level2)`: if a profile file already
exists with the same level2, the row is skipped unless --regenerate is passed.
zhao_001 is always preserved and never overwritten.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = REPO_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from heated_topics_v3.llm_client import call_llm  # noqa: E402
from heated_topics_v3.persona_intake import register_persona  # noqa: E402

DEFAULT_XLSX = Path("人设数据.xlsx")
DEFAULT_OUT = REPO_ROOT / "config" / "profiles"
DEFAULT_HDR = ("一级赛道", "二级赛道", "人设")

# Hand-curated profiles that must never be overwritten.
PRESERVED_USER_IDS: set[str] = {"zhao_001"}


def _iter_rows(xlsx: Path):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    col = {str(name).strip(): idx for idx, name in enumerate(header)}
    for row in rows[1:]:
        level1 = row[col.get(DEFAULT_HDR[0], -1)]
        level2 = row[col.get(DEFAULT_HDR[1], -1)]
        persona_text = row[col.get(DEFAULT_HDR[2], -1)]
        if not (level1 and level2 and persona_text):
            continue
        yield str(level1).strip(), str(level2).strip(), str(persona_text).strip()


def main(
    *,
    xlsx: Path = DEFAULT_XLSX,
    out_dir: Path = DEFAULT_OUT,
    use_llm: bool = True,
    regenerate: bool = False,
) -> int:
    if not xlsx.exists():
        print(f"xlsx not found: {xlsx}", file=sys.stderr)
        return 1

    cache_dir = REPO_ROOT / "cache" / "core_keywords"
    llm = None if not use_llm else lambda p, s, **kw: call_llm(p, system=s, **kw)

    written, skipped, errors = [], [], []

    for level1, level2, persona_text in _iter_rows(xlsx):
        # Check if level2 already has a profile (dedup by level2)
        existing_match = None
        slug = _slug_for(level2)
        for existing_file in out_dir.glob(f"{slug}_*.json"):
            try:
                import json

                payload = json.loads(existing_file.read_text(encoding="utf-8-sig"))
            except Exception:
                continue
            if payload.get("level2") == level2:
                existing_match = existing_file
                break

        if existing_match and existing_match.stem in PRESERVED_USER_IDS:
            skipped.append((level2, f"preserved ({existing_match.name})"))
            continue
        if existing_match and not regenerate:
            skipped.append((level2, "already exists"))
            continue

        try:
            if existing_match and regenerate:
                # Regenerate: reuse existing user_id, re-run LLM structure + keywords
                user_id = existing_match.stem
                result = _regenerate_profile(
                    existing_match,
                    level1,
                    level2,
                    persona_text,
                    cache_dir,
                    llm,
                )
                written.append((user_id, result))
                print(f"  ~ {user_id} ({level2}) [regenerated]")
            else:
                # New registration
                result = register_persona(
                    level1=level1,
                    level2=level2,
                    persona_text=persona_text,
                    profiles_dir=out_dir,
                    keyword_cache_dir=cache_dir,
                    use_llm=use_llm,
                    structurer_llm=llm,
                    keyword_llm=llm,
                )
                written.append((result.user_id, result.profile_path))
                print(f"  + {result.user_id} ({level2})")
        except Exception as exc:
            errors.append((level2, str(exc)))
            print(f"  ! {level2}: {exc}", file=sys.stderr)

    print(f"\nwrote {len(written)}, skipped {len(skipped)}, errors {len(errors)}")
    if skipped:
        print("skipped rows:")
        for level2, reason in skipped:
            print(f"  {level2} -> {reason}")
    if errors:
        print("errors:")
        for level2, err in errors:
            print(f"  {level2}: {err}")
    return 0


def _regenerate_profile(
    profile_path: Path,
    level1: str,
    level2: str,
    persona_text: str,
    cache_dir: Path,
    llm: callable | None,
) -> Path:
    """Overwrite existing profile with re-structured persona + fresh LLM keywords."""
    import json

    from heated_topics_v3.llm_keywords import extract_persona_keywords
    from heated_topics_v3.persona_structurer import structure_persona
    from heated_topics_v3.profile_loader import load_persona_profile

    # Re-structure personal fields
    try:
        personal = structure_persona(level1, level2, persona_text, llm=llm)
    except Exception:
        from heated_topics_v3.persona_structurer import _split_heuristic
        from heated_topics_v3.contracts import PersonaPersonal

        role, subject, scenarios, value = _split_heuristic(persona_text)
        if not role:
            role = level1
        if not subject:
            subject = level2
        if not value:
            value = f"围绕{level2}分享"
        personal = PersonaPersonal(
            role=role,
            subject=subject,
            scenarios=tuple(scenarios[:6]) if scenarios else (),
            value=value,
        )

    # Build minimal profile for keyword extraction
    temp_profile_path = profile_path.parent / f"_temp_{profile_path.stem}.json"
    import json as jsonmod

    temp_payload = {
        "user_id": profile_path.stem,
        "level1": level1,
        "level2": level2,
        "personal": {
            "role": personal.role,
            "subject": personal.subject,
            "scenarios": list(personal.scenarios),
            "value": personal.value,
        },
        "core_keywords": [],
    }
    temp_profile_path.write_text(
        jsonmod.dumps(temp_payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    try:
        profile = load_persona_profile(temp_profile_path)
        extraction = extract_persona_keywords(
            profile,
            cache_dir=cache_dir,
            llm=llm,
            allow_llm=llm is not None,
        )
        new_keywords = [k.keyword for k in extraction.keywords]
    finally:
        temp_profile_path.unlink(missing_ok=True)

    # Write back to original file
    payload = jsonmod.loads(profile_path.read_text(encoding="utf-8-sig"))
    payload["level1"] = level1
    payload["level2"] = level2
    payload["personal"] = {
        "role": personal.role,
        "subject": personal.subject,
        "scenarios": list(personal.scenarios),
        "value": personal.value,
    }
    payload["core_keywords"] = new_keywords

    profile_path.write_text(
        jsonmod.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return profile_path


def _slug_for(level2: str) -> str:
    from heated_topics_v3.persona_slugs import slug_for_level2

    return slug_for_level2(level2)


def _cli() -> int:
    p = argparse.ArgumentParser(description="Build v2 persona JSON files from an xlsx.")
    p.add_argument("--input", dest="xlsx", type=Path, default=DEFAULT_XLSX)
    p.add_argument("--out", dest="out_dir", type=Path, default=DEFAULT_OUT)
    p.add_argument("--use-llm", action="store_true", default=True)
    p.add_argument("--no-llm", action="store_true")
    p.add_argument("--regenerate", action="store_true")
    args = p.parse_args()
    if args.no_llm:
        use_llm = False
    else:
        use_llm = args.use_llm
    return main(
        xlsx=args.xlsx,
        out_dir=args.out_dir,
        use_llm=use_llm,
        regenerate=args.regenerate,
    )


if __name__ == "__main__":
    raise SystemExit(_cli())
